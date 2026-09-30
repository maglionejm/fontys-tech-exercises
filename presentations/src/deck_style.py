"""Shared design system for the M2 course decks.

Modern white + blue. 16:9. No author metadata (decks carry no personal details).
Every builder imports from here so all four decks look like one family.
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

# ---------------------------------------------------------------- tokens
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
NAVY = RGBColor(0x0A, 0x25, 0x40)   # headings, section slides
BLUE = RGBColor(0x1E, 0x5E, 0xF3)   # primary accent
SKY = RGBColor(0xB9, 0xD0, 0xFA)    # thin rules, soft accents
PANEL = RGBColor(0xEE, 0xF3, 0xFD)  # light panel fill
INK = RGBColor(0x20, 0x2B, 0x38)    # body text
GRAY = RGBColor(0x5B, 0x67, 0x74)   # secondary text
FONT = "Arial"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.6)
CONTENT_W = Inches(12.13)


# ---------------------------------------------------------------- basics
def new_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    shape.shadow.inherit = False


def _rect(slide, left, t, w, h, color):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, t, w, h)
    _fill(shp, color)
    return shp


def _box(slide, left, t, w, h):
    box = slide.shapes.add_textbox(left, t, w, h)
    box.text_frame.word_wrap = True
    return box


def _para(tf, text, size, color, bold=False, first=False, align=PP_ALIGN.LEFT,
          space_after=6, bullet=False, level=0):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(space_after)
    p.level = level
    run = p.add_run()
    run.text = text
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.color.rgb = color
    f.bold = bold
    if bullet:
        _bullet_char(p, level)
    else:
        _no_bullet(p)
    return p


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


def _bullet_char(p, level=0):
    pPr = p._p.get_or_add_pPr()
    pPr.set("indent", "-228600")
    pPr.set("marL", str(228600 + level * 285750))
    color_el = pPr.makeelement(qn("a:buClr"), {})
    srgb = pPr.makeelement(qn("a:srgbClr"), {"val": "1E5EF3"})
    color_el.append(srgb)
    char = pPr.makeelement(qn("a:buChar"), {"char": "—" if level == 0 else "·"})
    pPr.append(color_el)
    pPr.append(char)


def _header(slide, kicker, title):
    """Standard content-slide header: blue kicker, navy title, thin rule."""
    _rect(slide, MARGIN, Inches(0.52), Inches(0.5), Inches(0.06), BLUE)
    if kicker:
        kb = _box(slide, MARGIN, Inches(0.66), CONTENT_W, Inches(0.3))
        _para(kb.text_frame, kicker.upper(), 11, BLUE, bold=True, first=True)
    tb = _box(slide, MARGIN, Inches(0.95), CONTENT_W, Inches(0.85))
    _para(tb.text_frame, title, 26, NAVY, bold=True, first=True, space_after=0)
    _rect(slide, MARGIN, Inches(1.82), CONTENT_W, Inches(0.014), SKY)
    return Inches(2.05)  # y where content starts


# ---------------------------------------------------------------- slides
def title_slide(prs, kicker, title, subtitle, course="Machine Learning - Module 2"):
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    _rect(s, 0, 0, Inches(0.18), SLIDE_H, BLUE)
    kb = _box(s, Inches(1.0), Inches(2.15), Inches(11), Inches(0.4))
    _para(kb.text_frame, kicker.upper(), 14, BLUE, bold=True, first=True)
    tb = _box(s, Inches(1.0), Inches(2.6), Inches(11.3), Inches(1.9))
    _para(tb.text_frame, title, 44, NAVY, bold=True, first=True, space_after=0)
    sb = _box(s, Inches(1.0), Inches(4.5), Inches(10.5), Inches(0.9))
    _para(sb.text_frame, subtitle, 18, GRAY, first=True)
    cb = _box(s, Inches(1.0), Inches(6.6), Inches(10), Inches(0.4))
    _para(cb.text_frame, course, 12, GRAY, first=True)
    return s


def section_slide(prs, number, title, subtitle=""):
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    nb = _box(s, Inches(0.9), Inches(1.6), Inches(4), Inches(2.2))
    _para(nb.text_frame, number, 110, BLUE, bold=True, first=True)
    tb = _box(s, Inches(0.95), Inches(3.9), Inches(11.4), Inches(1.4))
    _para(tb.text_frame, title, 34, WHITE, bold=True, first=True)
    if subtitle:
        sb = _box(s, Inches(0.95), Inches(5.15), Inches(11), Inches(1.0))
        _para(sb.text_frame, subtitle, 16, SKY, first=True)
    return s


def bullets_slide(prs, title, bullets, kicker="", note=None):
    """bullets: list of str, or (str, [sub-lines]) tuples."""
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    y = _header(s, kicker, title)
    bb = _box(s, MARGIN, y, CONTENT_W, Inches(4.6))
    tf = bb.text_frame
    first = True
    for item in bullets:
        text, subs = item if isinstance(item, tuple) else (item, [])
        _para(tf, text, 16, INK, first=first, bullet=True, space_after=8)
        first = False
        for sub in subs:
            _para(tf, sub, 13.5, GRAY, bullet=True, level=1, space_after=6)
    if note:
        nb = _box(s, MARGIN, Inches(6.55), CONTENT_W, Inches(0.55))
        _para(nb.text_frame, note, 12, BLUE, first=True)
    return s


def _panel(slide, left, t, w, h, heading, lines, fill=PANEL, head_color=NAVY):
    _rect(slide, left, t, w, h, fill)
    hb = _box(slide, left + Inches(0.25), t + Inches(0.18), w - Inches(0.5), Inches(0.4))
    _para(hb.text_frame, heading, 15, head_color, bold=True, first=True)
    bb = _box(slide, left + Inches(0.25), t + Inches(0.62), w - Inches(0.5), h - Inches(0.8))
    tf = bb.text_frame
    first = True
    for line in lines:
        _para(tf, line, 13, INK, first=first, bullet=True, space_after=6)
        first = False


def two_col_slide(prs, title, left, right, kicker="", note=None):
    """left/right: (heading, [lines])."""
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    y = _header(s, kicker, title)
    h = Inches(4.35) if note else Inches(4.7)
    col_w = Inches(5.9)
    _panel(s, MARGIN, y, col_w, h, left[0], left[1])
    _panel(s, MARGIN + col_w + Inches(0.33), y, col_w, h, right[0], right[1])
    if note:
        nb = _box(s, MARGIN, Inches(6.62), CONTENT_W, Inches(0.55))
        _para(nb.text_frame, note, 12, BLUE, first=True)
    return s


def image_slide(prs, title, img_path, kicker="", bullets=None, caption=""):
    """Image on the right, optional bullets on the left."""
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    y = _header(s, kicker, title)
    if bullets:
        bb = _box(s, MARGIN, y, Inches(4.7), Inches(4.6))
        tf = bb.text_frame
        first = True
        for b in bullets:
            _para(tf, b, 14.5, INK, first=first, bullet=True, space_after=8)
            first = False
        img_l, img_w = Inches(5.6), Inches(7.1)
    else:
        img_l, img_w = Inches(1.7), Inches(9.9)
    pic = s.shapes.add_picture(img_path, img_l, y, width=img_w)
    max_h = Inches(4.5)
    if pic.height > max_h:  # keep inside the content band
        scale = max_h / pic.height
        pic.height = int(pic.height * scale)
        pic.width = int(pic.width * scale)
        pic.left = img_l + int((img_w - pic.width) / 2)
    if caption:
        cb = _box(s, img_l, y + max_h + Inches(0.08), img_w, Inches(0.4))
        _para(cb.text_frame, caption, 11, GRAY, first=True, align=PP_ALIGN.CENTER)
    return s


def table_slide(prs, title, headers, rows, kicker="", note=None, col_widths=None):
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    y = _header(s, kicker, title)
    n_rows, n_cols = len(rows) + 1, len(headers)
    h = Inches(0.5) * n_rows
    gfx = s.shapes.add_table(n_rows, n_cols, MARGIN, y, CONTENT_W, min(h, Inches(4.5)))
    table = gfx.table
    if col_widths:
        total = sum(col_widths)
        for i, wgt in enumerate(col_widths):
            table.columns[i].width = int(CONTENT_W * wgt / total)
    for c, head in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        cell.text_frame.word_wrap = True
        _para(cell.text_frame, head, 13, WHITE, bold=True, first=True, space_after=0)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else PANEL
            cell.text_frame.word_wrap = True
            _para(cell.text_frame, str(val), 12.5, INK, first=True, space_after=0)
    if note:
        nb = _box(s, MARGIN, Inches(6.62), CONTENT_W, Inches(0.55))
        _para(nb.text_frame, note, 12, BLUE, first=True)
    return s


def big_number_slide(prs, title, number, label, foot="", kicker=""):
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    _header(s, kicker, title)
    nb = _box(s, Inches(1.5), Inches(2.5), Inches(10.3), Inches(2.2))
    _para(nb.text_frame, number, 96, BLUE, bold=True, first=True, align=PP_ALIGN.CENTER)
    lb = _box(s, Inches(1.5), Inches(4.85), Inches(10.3), Inches(0.8))
    _para(lb.text_frame, label, 18, NAVY, first=True, align=PP_ALIGN.CENTER)
    if foot:
        fb = _box(s, Inches(1.5), Inches(5.75), Inches(10.3), Inches(0.7))
        _para(fb.text_frame, foot, 12.5, GRAY, first=True, align=PP_ALIGN.CENTER)
    return s


def quote_slide(prs, text, attribution):
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, WHITE)
    _rect(s, Inches(1.0), Inches(1.7), Inches(0.08), Inches(3.0), BLUE)
    qb = _box(s, Inches(1.5), Inches(1.8), Inches(10.3), Inches(2.8))
    _para(qb.text_frame, text, 26, NAVY, first=True)
    ab = _box(s, Inches(1.5), Inches(5.0), Inches(10.3), Inches(0.6))
    _para(ab.text_frame, attribution, 14, GRAY, first=True)
    return s


def close_slide(prs, title, lines, course="Machine Learning - Module 2"):
    s = _blank(prs)
    _rect(s, 0, 0, SLIDE_W, SLIDE_H, NAVY)
    tb = _box(s, Inches(1.0), Inches(1.4), Inches(11.3), Inches(1.2))
    _para(tb.text_frame, title, 34, WHITE, bold=True, first=True)
    bb = _box(s, Inches(1.0), Inches(2.9), Inches(11.0), Inches(3.2))
    tf = bb.text_frame
    first = True
    for line in lines:
        _para(tf, line, 16, SKY, first=first, bullet=True, space_after=10)
        first = False
    cb = _box(s, Inches(1.0), Inches(6.6), Inches(10), Inches(0.4))
    _para(cb.text_frame, course, 12, SKY, first=True)
    return s


# ---------------------------------------------------------------- finish
def save_deck(prs, path, footer_text, author="M2 Machine Learning course"):
    """Add footers/page numbers (skipping title slide), scrub metadata, save.

    author is the generic course name written into the file's properties in
    place of any personal name (the decks carry no personal details).
    """
    for idx, slide in enumerate(prs.slides, start=1):
        if idx == 1:
            continue
        # detect dark slides by our own convention: full-bleed NAVY rect first
        is_dark = False
        try:
            shp = slide.shapes[0]
            if shp.width == SLIDE_W and shp.fill.type is not None:
                is_dark = shp.fill.fore_color.rgb == NAVY
        except Exception:
            pass
        color = SKY if is_dark else GRAY
        fb = _box(slide, MARGIN, Inches(7.08), Inches(8), Inches(0.32))
        _para(fb.text_frame, footer_text, 9, color, first=True)
        nb = _box(slide, Inches(12.35), Inches(7.08), Inches(0.6), Inches(0.32))
        _para(nb.text_frame, str(idx), 9, color, first=True, align=PP_ALIGN.RIGHT)
    cp = prs.core_properties
    cp.author = author
    cp.last_modified_by = author
    cp.comments = ""
    cp.title = footer_text
    prs.save(path)


# ---------------------------------------------------------------- charts
def mpl_theme():
    """Apply the deck look to matplotlib. Returns hex palette dict."""
    import matplotlib as mpl
    pal = {"navy": "#0A2540", "blue": "#1E5EF3", "sky": "#B9D0FA",
           "panel": "#EEF3FD", "ink": "#202B38", "gray": "#5B6774"}
    mpl.rcParams.update({
        "font.family": "Arial", "font.size": 13,
        "text.color": pal["ink"], "axes.edgecolor": pal["gray"],
        "axes.labelcolor": pal["ink"], "axes.titlesize": 15,
        "axes.titleweight": "bold", "axes.titlecolor": pal["navy"],
        "axes.spines.top": False, "axes.spines.right": False,
        "xtick.color": pal["gray"], "ytick.color": pal["gray"],
        "figure.facecolor": "white", "axes.facecolor": "white",
        "savefig.dpi": 200, "savefig.bbox": "tight",
    })
    return pal
